import streamlit as st
import pdfplumber
from pypdf import PdfReader, PdfWriter
from pdf2image import convert_from_bytes
import pytesseract
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import re
import zipfile
import io

st.set_page_config(page_title="Invoice Splitter", page_icon="📄", layout="centered")

st.markdown("""
<style>
    @import url('https://fonts.googleapis.com/css2?family=Barlow+Condensed:wght@400;600;700&family=Barlow:wght@400;500&display=swap');
    html, body, [class*="css"] { font-family: 'Barlow', sans-serif; }
    h1, h2, h3 { font-family: 'Barlow Condensed', sans-serif !important; letter-spacing: 0.5px; }
    .stButton > button {
        background-color: #003399; color: white;
        font-family: 'Barlow Condensed', sans-serif;
        font-size: 18px; font-weight: 600; letter-spacing: 1px;
        border: none; border-radius: 4px; padding: 0.6em 2em; width: 100%;
    }
    .stButton > button:hover { background-color: #002277; color: white; }
    .download-bar {
        background: #f0f4ff; border: 1px solid #003399;
        border-radius: 6px; padding: 14px 18px; margin-bottom: 12px;
    }
    .download-bar-title { font-family: 'Barlow Condensed', sans-serif; font-size: 15px; font-weight: 600; color: #003399; margin-bottom: 10px; }
    .stat-box { background: white; border-left: 4px solid #003399; border-radius: 4px; padding: 16px 20px; margin: 8px 0; box-shadow: 0 1px 4px rgba(0,0,0,0.08); }
    .stat-number { font-family: 'Barlow Condensed', sans-serif; font-size: 36px; font-weight: 700; color: #003399; line-height: 1; }
    .stat-label  { font-size: 13px; color: #666; margin-top: 4px; text-transform: uppercase; letter-spacing: 0.5px; }
    .invoice-row { background: white; border-radius: 4px; padding: 10px 14px; margin: 4px 0; font-size: 14px; box-shadow: 0 1px 3px rgba(0,0,0,0.06); }
    .badge-new   { background:#e6f0ff; color:#003399; padding:2px 8px; border-radius:3px; font-size:12px; font-weight:600; }
    .badge-used  { background:#fff3e6; color:#cc6600; padding:2px 8px; border-radius:3px; font-size:12px; font-weight:600; }
    .badge-clean { background:#e6ffe6; color:#006600; padding:2px 8px; border-radius:3px; font-size:12px; font-weight:600; }
    .badge-other { background:#f0f0f0; color:#555;    padding:2px 8px; border-radius:3px; font-size:12px; font-weight:600; }
    .ocr-note    { font-size:11px; color:#888; margin-left:6px; }
    .upload-hint { text-align:center; color:#888; font-size:13px; margin-top:-8px; margin-bottom:16px; }
</style>
""", unsafe_allow_html=True)

st.markdown("# 📄 Invoice Splitter")
st.markdown("Upload your combined invoice PDF — splits into individual files and generates an Excel summary.")
st.divider()

# ── Helpers ───────────────────────────────────────────────────

def ocr_page_image(pdf_bytes, page_num):
    try:
        images = convert_from_bytes(pdf_bytes, first_page=page_num + 1, last_page=page_num + 1, dpi=300)
        if images:
            return pytesseract.image_to_string(images[0])
    except Exception as e:
        st.warning(f"OCR error on page {page_num + 1}: {e}")
    return ""

def get_page_text(pdf_bytes, plumber_page, page_num):
    text = plumber_page.extract_text() or ""
    if len(text.strip()) < 50:
        text = ocr_page_image(pdf_bytes, page_num)
        return text, True
    return text, False

def classify_service(text):
    t = text.lower()
    if "used detail"        in t: return "Used Detail"
    if "new detail"         in t: return "New Detail"
    if "clean for delivery" in t: return "Clean For Delivery"
    if "freshen up"         in t: return "Freshen Up"
    if "showroom wash"      in t: return "Showroom Wash"
    return "Other"

def get_badge(service):
    s = service.lower()
    if "new"    in s: return "new"
    if "used"   in s: return "used"
    if "clean" in s or "freshen" in s or "showroom" in s: return "clean"
    return "other"

def extract_invoice_info(text):
    inv = re.search(r'Invoice\s*(?:Number|#|No\.?)[:\s]*([A-Z]{2,5}-\d+)', text, re.IGNORECASE)
    invoice_num = inv.group(1).upper() if inv else None
    stock = re.search(r'Stk[:\s#]*(\S+)', text, re.IGNORECASE)
    stock_num = stock.group(1).strip() if stock else "UNKNOWN"
    date_m = re.search(r'(?:Date|Dated)[:\s]*(\d{1,2}/\d{1,2}/\d{2,4})', text, re.IGNORECASE)
    date = date_m.group(1) if date_m else "UNKNOWN"
    vehicle = "UNKNOWN"
    for p in [
        r'(\d{4}\s+[A-Z][a-z]+\s+[^\n\r]{3,40}?)(?:\n|\r|$)',
        r'(?:VIN|Vin)[^\n]*:\s*[A-Z0-9]{17}[:\s]+(.+?)(?:\n|$)',
        r'Clean:\s*\S+:\s*(\d{4}\s+[^\n]+)',
    ]:
        m = re.search(p, text)
        if m:
            v = m.group(1).strip()
            if re.match(r'^\d{4}', v):
                vehicle = v[:50]
                break
    service = classify_service(text)
    amounts = re.findall(r'\$?\s*(\d{1,6}\.\d{2})', text)
    total   = float(amounts[-1]) if amounts else 0.0
    return invoice_num, stock_num, date, vehicle, service, total

def safe_filename(s):
    return re.sub(r'[\\/:*?"<>|]', '', s).strip()

def build_excel(records):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Invoice Summary"
    thin   = Side(style='thin', color='CCCCCC')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    headers = ['Invoice #', 'Stock #', 'Date', 'Vehicle', 'Service Type', 'Total']
    widths  = [14, 12, 12, 35, 22, 12]
    for col, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font      = Font(name='Calibri', bold=True, size=11, color='FFFFFF')
        cell.fill      = PatternFill('solid', start_color='003399')
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border    = border
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = w
    ws.row_dimensions[1].height = 22
    type_colors = {'New':'E2EFDA','Used':'FCE4D6','Clean':'DDEBF7','Freshen':'FFF2CC','Other':'F2F2F2'}
    for row_idx, rec in enumerate(records, 2):
        ck       = next((k for k in type_colors if k.lower() in rec['service'].lower()), 'Other')
        row_fill = PatternFill('solid', start_color=type_colors[ck])
        data     = [rec['invoice'], rec['stock'], rec['date'], rec['vehicle'], rec['service'], rec['total']]
        for col, val in enumerate(data, 1):
            cell           = ws.cell(row=row_idx, column=col, value=val)
            cell.font      = Font(name='Calibri', size=10)
            cell.fill      = row_fill
            cell.border    = border
            cell.alignment = Alignment(horizontal='center' if col not in (1,4) else 'left', vertical='center')
            if col == 6: cell.number_format = '$#,##0.00'
    total_row = len(records) + 2
    lbl = ws.cell(row=total_row, column=5, value='GRAND TOTAL')
    lbl.font = Font(name='Calibri', bold=True, size=10)
    lbl.alignment = Alignment(horizontal='right')
    tot = ws.cell(row=total_row, column=6, value=f'=SUM(F2:F{len(records)+1})')
    tot.font          = Font(name='Calibri', bold=True, size=10)
    tot.number_format = '$#,##0.00'
    tot.border        = Border(top=Side(style='double', color='000000'))
    ws.freeze_panes   = 'A2'
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf

# ── Session state ─────────────────────────────────────────────
for key, default in [('inv_ready',False),('inv_zip',None),('inv_excel',None),('inv_records',[])]:
    if key not in st.session_state:
        st.session_state[key] = default

# ── Top download bar ──────────────────────────────────────────
if st.session_state.inv_ready:
    st.markdown('<div class="download-bar"><div class="download-bar-title">✅ Files ready — download below</div></div>', unsafe_allow_html=True)
    dl1, dl2 = st.columns(2)
    with dl1:
        st.download_button("📦 Download All PDFs (ZIP)", data=st.session_state.inv_zip,
                           file_name="invoices_split.zip", mime="application/zip",
                           use_container_width=True, key="top_zip")
    with dl2:
        st.download_button("📊 Download Excel Summary", data=st.session_state.inv_excel,
                           file_name="invoice_summary.xlsx",
                           mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                           use_container_width=True, key="top_excel")
    st.divider()

# ── Upload ────────────────────────────────────────────────────
uploaded_file = st.file_uploader("Upload Combined Invoice PDF", type=["pdf"])
st.markdown('<p class="upload-hint">Your file is never stored — processing happens in memory only.</p>', unsafe_allow_html=True)

if uploaded_file:
    st.success(f"✅ **{uploaded_file.name}** uploaded")
    if st.button("⚙️ SPLIT INVOICES"):
        st.session_state.inv_ready = False
        pdf_bytes   = uploaded_file.read()
        total_pages = len(PdfReader(io.BytesIO(pdf_bytes)).pages)
        records     = []
        progress    = st.progress(0, text="Processing pages…")
        status_ph   = st.empty()
        with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
            for i, page in enumerate(pdf.pages):
                progress.progress((i+1)/total_pages, text=f"Processing page {i+1} of {total_pages}…")
                text, used_ocr = get_page_text(pdf_bytes, page, i)
                if used_ocr:
                    status_ph.info(f"📷 Page {i+1}: scanned page — using OCR")
                inv_num, stock, date, vehicle, service, total = extract_invoice_info(text)
                if not inv_num: inv_num = f"PAGE-{i+1}"
                records.append({'page_idx':i,'invoice':inv_num,'stock':stock,'date':date,
                                'vehicle':vehicle,'service':service,'total':total,'ocr':used_ocr})
        progress.empty(); status_ph.empty()
        with st.spinner("Packaging PDFs…"):
            zip_buf = io.BytesIO()
            reader2 = PdfReader(io.BytesIO(pdf_bytes))
            with zipfile.ZipFile(zip_buf, 'w', zipfile.ZIP_DEFLATED) as zf:
                for rec in records:
                    writer  = PdfWriter()
                    writer.add_page(reader2.pages[rec['page_idx']])
                    fname   = safe_filename(f"{rec['invoice']} - {rec['service']} - {rec['vehicle']}.pdf")
                    pdf_out = io.BytesIO()
                    writer.write(pdf_out)
                    zf.writestr(fname, pdf_out.getvalue())
            zip_buf.seek(0)
        excel_buf = build_excel(records)
        st.session_state.inv_ready   = True
        st.session_state.inv_zip     = zip_buf.getvalue()
        st.session_state.inv_excel   = excel_buf.getvalue()
        st.session_state.inv_records = records
        st.rerun()

# ── Results ───────────────────────────────────────────────────
if st.session_state.inv_ready and st.session_state.inv_records:
    records     = st.session_state.inv_records
    new_count   = sum(1 for r in records if 'new'  in r['service'].lower())
    used_count  = sum(1 for r in records if 'used' in r['service'].lower())
    grand_total = sum(r['total'] for r in records)
    ocr_pages   = [r['page_idx']+1 for r in records if r['ocr']]
    st.markdown("## Results")
    c1, c2, c3 = st.columns(3)
    with c1: st.markdown(f'<div class="stat-box"><div class="stat-number">{len(records)}</div><div class="stat-label">Total Invoices</div></div>', unsafe_allow_html=True)
    with c2: st.markdown(f'<div class="stat-box"><div class="stat-number">{new_count} / {used_count}</div><div class="stat-label">New / Used</div></div>', unsafe_allow_html=True)
    with c3: st.markdown(f'<div class="stat-box"><div class="stat-number">${grand_total:,.2f}</div><div class="stat-label">Grand Total</div></div>', unsafe_allow_html=True)
    if ocr_pages:
        st.info(f"📷 OCR used on {len(ocr_pages)} page(s): {', '.join(map(str, ocr_pages))}")
    st.markdown("### Invoice Breakdown")
    for rec in records:
        badge   = get_badge(rec['service'])
        ocr_tag = ' <span class="ocr-note">📷 OCR</span>' if rec['ocr'] else ''
        st.markdown(
            f'<div class="invoice-row"><b>{rec["invoice"]}</b>&nbsp;&nbsp;'
            f'<span class="badge-{badge}">{rec["service"]}</span>{ocr_tag}&nbsp;&nbsp;'
            f'<span style="color:#555">{rec["vehicle"]}</span>'
            f'<span style="float:right;font-weight:600">${rec["total"]:,.2f}</span></div>',
            unsafe_allow_html=True)
    st.divider()
    st.caption("Each PDF is named: Invoice # · Service Type · Vehicle")
